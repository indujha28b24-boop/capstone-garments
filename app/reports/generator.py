import os
from io import BytesIO
from datetime import datetime
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models import Order, Inventory, Worker, Machine, ProductionTracking

class ReportGenerator:

    @staticmethod
    def generate_production_pdf():
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        # Title Header
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#1e293b'),
            alignment=0
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#64748b')
        )

        elements.append(Paragraph("STITCHTECH ERP - GARMENT PRODUCTION REPORT", title_style))
        elements.append(Paragraph(f"Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | Commercial Factory Operations", subtitle_style))
        elements.append(Spacer(1, 15))

        # Table Header & Data
        headers = ['Order ID', 'Garment Type', 'Quantity', 'Stage', 'Progress', 'Deadline', 'Status']
        data = [headers]

        orders = Order.query.order_by(Order.id.desc()).all()
        for ord in orders:
            data.append([
                f"#{ord.id}",
                ord.garment_type,
                f"{ord.quantity:,}",
                ord.current_stage,
                f"{ord.progress_percentage}%",
                ord.deadline.strftime('%Y-%m-%d'),
                ord.status
            ])

        t = Table(data, colWidths=[60, 90, 60, 100, 60, 80, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0284c7')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (2, 1), (2, -1), 'RIGHT'),
            ('ALIGN', (4, 1), (4, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
        ]))

        elements.append(t)
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_production_excel():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Production Schedule"

        # Styling
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        headers = ['Order ID', 'Customer ID', 'Garment Type', 'Quantity', 'Current Stage', 'Progress (%)', 'Order Date', 'Deadline', 'Status']
        ws.append(headers)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        orders = Order.query.order_by(Order.id.desc()).all()
        for ord in orders:
            row = [
                f"ORD-{ord.id:04d}",
                f"CUST-{ord.customer_id:03d}",
                ord.garment_type,
                ord.quantity,
                ord.current_stage,
                ord.progress_percentage,
                ord.order_date.strftime('%Y-%m-%d'),
                ord.deadline.strftime('%Y-%m-%d'),
                ord.status
            ]
            ws.append(row)
            row_idx = ws.max_row
            
            for col_idx in range(1, len(row) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx in [4, 6]:
                    cell.alignment = align_right
                elif col_idx in [1, 2, 7, 8, 9]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_inventory_pdf():
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('InvTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=20, textColor=colors.HexColor('#0f172a'))
        elements.append(Paragraph("STITCHTECH ERP - INVENTORY AUDIT REPORT", title_style))
        elements.append(Paragraph(f"Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}", styles['Normal']))
        elements.append(Spacer(1, 15))

        headers = ['ID', 'Item Name', 'Category', 'Quantity', 'Unit', 'Min Stock', 'Stock Status']
        data = [headers]

        inventory_items = Inventory.query.all()
        for item in inventory_items:
            status = 'LOW STOCK' if item.quantity <= item.minimum_stock else 'HEALTHY'
            data.append([
                f"INV-{item.id:03d}",
                item.item_name,
                item.category,
                f"{item.quantity:,}",
                item.unit,
                f"{item.minimum_stock:,}",
                status
            ])

        t = Table(data, colWidths=[60, 130, 90, 60, 50, 60, 80])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f766e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0fdf4')]),
        ]))

        elements.append(t)
        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_inventory_excel():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory Audit"

        header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        headers = ['Item ID', 'Item Name', 'Category', 'Quantity', 'Unit', 'Minimum Stock', 'Stock Status']
        ws.append(headers)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        inventory_items = Inventory.query.all()
        for item in inventory_items:
            status = 'LOW STOCK' if item.quantity <= item.minimum_stock else 'HEALTHY'
            ws.append([
                f"INV-{item.id:03d}",
                item.item_name,
                item.category,
                item.quantity,
                item.unit,
                item.minimum_stock,
                status
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
